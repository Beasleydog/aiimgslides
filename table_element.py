import random

from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

import config
from element_model import Element
from utils import add_wrapped_text, draw_rect_outline, ppt_color, rand_color, rand_text


TABLE_DASHES = ["solid", "dash", "dashDot", "lgDash", "sysDot"]


def hex_color(rgb):
    return "%02X%02X%02X" % rgb[:3]


def soften(rgb, amount=0.82):
    return tuple(int(c + (255 - c) * amount) for c in rgb[:3])


def set_cell_border(cell, color, width_pt, dash):
    tc_pr = cell._tc.get_or_add_tcPr()
    for side in ["lnL", "lnR", "lnT", "lnB"]:
        existing = tc_pr.find(qn(f"a:{side}"))
        if existing is not None:
            tc_pr.remove(existing)
        line = OxmlElement(f"a:{side}")
        line.set("w", str(int(width_pt * 12700)))

        solid_fill = OxmlElement("a:solidFill")
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", hex_color(color))
        solid_fill.append(srgb)
        line.append(solid_fill)

        preset_dash = OxmlElement("a:prstDash")
        preset_dash.set("val", dash)
        line.append(preset_dash)
        tc_pr.append(line)


def make_table(x, y, w, h):
    rows = random.randint(*config.TABLE_ROW_RANGE)
    cols = random.randint(*config.TABLE_COL_RANGE)
    header = rand_color()
    accent = rand_color()
    body = soften(accent, random.uniform(0.78, 0.94))
    border_color = rand_color()
    has_header = random.random() < config.TABLE_HEADER_ROW_PROBABILITY
    band_rows = random.random() < config.TABLE_BANDED_ROWS_PROBABILITY
    band_cols = random.random() < config.TABLE_BANDED_COLS_PROBABILITY
    return Element(
        "table",
        x,
        y,
        w,
        h,
        {
            "rows": rows,
            "cols": cols,
            "cells": [[rand_text(1, 2) for _ in range(cols)] for _ in range(rows)],
            "header": header,
            "body": body,
            "band": soften(accent, random.uniform(0.58, 0.76)),
            "border": border_color,
            "border_width": random.uniform(*config.TABLE_BORDER_WIDTH_RANGE),
            "border_dash": random.choice(TABLE_DASHES)
            if random.random() < config.TABLE_DASHED_BORDER_PROBABILITY
            else "solid",
            "font_size": random.randint(*config.TABLE_FONT_SIZE_RANGE),
            "text_color": rand_color(),
            "has_header": has_header,
            "band_rows": band_rows,
            "band_cols": band_cols,
        },
    )


def add_table_to_pptx(slide, el):
    table = slide.shapes.add_table(
        el.data["rows"],
        el.data["cols"],
        Inches(el.x),
        Inches(el.y),
        Inches(el.w),
        Inches(el.h),
    ).table
    for r in range(el.data["rows"]):
        for c in range(el.data["cols"]):
            cell = table.cell(r, c)
            cell.text = el.data["cells"][r][c]
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(el.data["font_size"])
            run.font.color.rgb = ppt_color(el.data["text_color"])
            run.font.bold = el.data["has_header"] and r == 0

            fill = el.data["body"]
            if el.data["has_header"] and r == 0:
                fill = el.data["header"]
            elif el.data["band_rows"] and r % 2 == 1:
                fill = el.data["band"]
            elif el.data["band_cols"] and c % 2 == 1:
                fill = el.data["band"]

            cell.fill.solid()
            cell.fill.fore_color.rgb = ppt_color(fill)
            set_cell_border(cell, el.data["border"], el.data["border_width"], el.data["border_dash"])


def add_table_to_png(draw, el, box):
    x1, y1, x2, y2 = box
    rows, cols = el.data["rows"], el.data["cols"]
    cw = (x2 - x1) / cols
    ch = (y2 - y1) / rows
    border = (*el.data["border"], 255)
    width = max(1, int(el.data["border_width"] * 1.6))
    dashed = el.data["border_dash"] != "solid"
    draw.rectangle(box, fill=(*el.data["body"], 220))
    draw_rect_outline(draw, box, border, width, dashed)
    for r in range(rows):
        for c in range(cols):
            cx1, cy1 = int(x1 + c * cw), int(y1 + r * ch)
            cx2, cy2 = int(x1 + (c + 1) * cw), int(y1 + (r + 1) * ch)
            fill = el.data["body"]
            if el.data["has_header"] and r == 0:
                fill = el.data["header"]
            elif el.data["band_rows"] and r % 2 == 1:
                fill = el.data["band"]
            elif el.data["band_cols"] and c % 2 == 1:
                fill = el.data["band"]
            cell_box = (cx1, cy1, cx2, cy2)
            draw.rectangle(cell_box, fill=(*fill, 205))
            draw_rect_outline(draw, cell_box, border, width, dashed)
            add_wrapped_text(
                draw,
                el.data["cells"][r][c],
                (cx1 + 5, cy1 + 4, cx2 - 5, cy2 - 4),
                el.data["text_color"],
                max(8, int(el.data["font_size"] * 1.35)),
                el.data["has_header"] and r == 0,
            )
